from pathlib import Path
from xml.etree.ElementTree import (
    Element,
    ElementTree,
    SubElement,
)

from docx import Document


OUTPUT_DIRECTORY = Path("data/incoming")
OUTPUT_DIRECTORY.mkdir(parents=True, exist_ok=True)


SKILLS = (
    "Python, JavaScript, React, Node.js, "
    "Express, MongoDB, Git"
)

PROJECT = (
    "Built a React dashboard and developed REST APIs "
    "using Node.js, Express, and MongoDB."
)

EXPERIENCE = (
    "Worked with a three-member engineering team and "
    "reduced API response time by 30 percent."
)


def create_docx_resume() -> None:
    document = Document()

    document.add_heading(
        "PRIYA VERMA",
        level=0,
    )
    document.add_paragraph(
        "Junior Full Stack Developer"
    )

    document.add_heading(
        "Technical Skills",
        level=1,
    )
    document.add_paragraph(SKILLS)

    document.add_heading(
        "Projects",
        level=1,
    )
    document.add_paragraph(
        PROJECT,
        style="List Bullet",
    )

    document.add_heading(
        "Experience",
        level=1,
    )
    document.add_paragraph(EXPERIENCE)

    document.add_heading(
        "Education",
        level=1,
    )

    table = document.add_table(
        rows=2,
        cols=2,
    )

    table.cell(0, 0).text = "Degree"
    table.cell(0, 1).text = "Institution"
    table.cell(1, 0).text = (
        "Bachelor of Technology in Computer Science"
    )
    table.cell(1, 1).text = "Example Institute"

    output_path = OUTPUT_DIRECTORY / "test_resume.docx"
    document.save(output_path)

    print(f"Created: {output_path}")


def create_txt_resume() -> None:
    text = f"""
ROHAN MEHTA
Backend Developer Intern

SKILLS
{SKILLS}

PROJECTS
{PROJECT}

EXPERIENCE
{EXPERIENCE}

EDUCATION
Bachelor of Technology in Computer Science
""".strip()

    output_path = OUTPUT_DIRECTORY / "test_resume.txt"

    # UTF-16 deliberately tests encoding detection.
    output_path.write_text(
        text,
        encoding="utf-16",
    )

    print(f"Created: {output_path}")


def create_xml_resume() -> None:
    resume = Element("Resume")

    identity = SubElement(resume, "Identity")
    SubElement(identity, "Name").text = "ANANYA GUPTA"
    SubElement(identity, "Title").text = (
        "Full Stack Developer Intern"
    )

    skills = SubElement(resume, "Skills")
    SubElement(skills, "SkillList").text = SKILLS

    projects = SubElement(resume, "Projects")
    SubElement(projects, "Project").text = PROJECT

    experience = SubElement(resume, "Experience")
    SubElement(experience, "Role").text = EXPERIENCE

    education = SubElement(resume, "Education")
    SubElement(education, "Degree").text = (
        "Bachelor of Technology in Computer Science"
    )

    output_path = OUTPUT_DIRECTORY / "test_resume.xml"

    ElementTree(resume).write(
        output_path,
        encoding="utf-8",
        xml_declaration=True,
    )

    print(f"Created: {output_path}")


if __name__ == "__main__":
    create_docx_resume()
    create_txt_resume()
    create_xml_resume()