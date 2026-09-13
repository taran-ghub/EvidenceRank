from __future__ import annotations

import argparse
from pathlib import Path

from defusedxml import ElementTree
from docx import Document
from docx.oxml.ns import qn

from src.extractor import build_page_result, extract_pdf
from src.schemas import ExtractedDocument, PageExtraction


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".xml"}


def failed_document(
    path: Path,
    warning: str,
) -> ExtractedDocument:
    return ExtractedDocument(
        filename=path.name,
        page_count=0,
        pages=[],
        full_text="",
        usable_page_count=0,
        extraction_success=False,
        warnings=[warning],
    )


def build_single_page_document(
    path: Path,
    text: str,
    method: str,
    extra_warnings: list[str] | None = None,
) -> ExtractedDocument:
    page = build_page_result(
        page_number=1,
        text=text,
        extraction_method=method,
    )

    if extra_warnings:
        page.warnings.extend(extra_warnings)

    warnings: list[str] = []

    if not page.is_usable:
        warnings.append(
            f"{method.upper()} extraction produced insufficient text."
        )

    return ExtractedDocument(
        filename=path.name,
        page_count=1,
        pages=[page],
        full_text=page.text,
        usable_page_count=1 if page.is_usable else 0,
        extraction_success=page.is_usable,
        warnings=warnings,
    )


def extract_docx(path: Path) -> ExtractedDocument:
    """
    Extract paragraphs, tables, and most text-box content from DOCX.

    Word files do not expose reliable page boundaries without rendering,
    so the extracted document is represented as one logical page.
    """
    try:
        document = Document(path)
        text_blocks: list[str] = []

        for paragraph_element in document.element.body.iter(
            qn("w:p")
        ):
            paragraph_text = "".join(
                text_node.text or ""
                for text_node in paragraph_element.iter(qn("w:t"))
            ).strip()

            if paragraph_text:
                text_blocks.append(paragraph_text)

        # Some resumes place information inside headers or footers.
        for section in document.sections:
            for container in (section.header, section.footer):
                for paragraph in container.paragraphs:
                    text = paragraph.text.strip()

                    if text:
                        text_blocks.append(text)

                for table in container.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            text = cell.text.strip()

                            if text:
                                text_blocks.append(text)

        # Remove consecutive duplicates that can appear in Word XML.
        deduplicated_blocks: list[str] = []

        for block in text_blocks:
            if (
                not deduplicated_blocks
                or block != deduplicated_blocks[-1]
            ):
                deduplicated_blocks.append(block)

        text = "\n".join(deduplicated_blocks)

        return build_single_page_document(
            path=path,
            text=text,
            method="docx",
            extra_warnings=[
                "DOCX page boundaries are unavailable; "
                "page number 1 is a logical document page."
            ],
        )

    except Exception as error:
        return failed_document(
            path,
            f"DOCX extraction failed: {error}",
        )


def decode_text_bytes(data: bytes) -> tuple[str, str]:
    """
    Detect common text encodings and retain the most readable result.
    """

    # UTF-16 files normally begin with one of these byte-order marks.
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        try:
            return data.decode("utf-16"), "utf-16"
        except UnicodeDecodeError:
            pass

    # Detect a UTF-8 byte-order mark.
    if data.startswith(b"\xef\xbb\xbf"):
        try:
            return data.decode("utf-8-sig"), "utf-8-sig"
        except UnicodeDecodeError:
            pass

    encodings = [
        "utf-8",
        "cp1252",
        "latin-1",
    ]

    candidates: list[tuple[int, str, str]] = []

    for encoding in encodings:
        try:
            text = data.decode(encoding)

            alphabetic_count = sum(
                character.isalpha()
                for character in text
            )

            replacement_count = text.count("�")
            null_count = text.count("\x00")

            control_character_count = sum(
                ord(character) < 32
                and character not in "\n\r\t"
                for character in text
            )

            score = (
                alphabetic_count
                - (replacement_count * 50)
                - (null_count * 100)
                - (control_character_count * 20)
            )

            candidates.append(
                (score, text, encoding)
            )

        except UnicodeDecodeError:
            continue

    if not candidates:
        return "", "unknown"

    _, best_text, best_encoding = max(
        candidates,
        key=lambda candidate: candidate[0],
    )

    return best_text, best_encoding
    """
    Try common text encodings and retain the most readable result.
    """
    encodings = [
        "utf-8-sig",
        "utf-16",
        "cp1252",
        "latin-1",
    ]

    candidates: list[tuple[int, str, str]] = []

    for encoding in encodings:
        try:
            text = data.decode(encoding)

            alphabetic_count = sum(
                character.isalpha()
                for character in text
            )

            replacement_count = text.count("�")

            score = (
                alphabetic_count
                - (replacement_count * 50)
            )

            candidates.append(
                (score, text, encoding)
            )

        except UnicodeDecodeError:
            continue

    if not candidates:
        return "", "unknown"

    _, best_text, best_encoding = max(
        candidates,
        key=lambda candidate: candidate[0],
    )

    return best_text, best_encoding


def extract_txt(path: Path) -> ExtractedDocument:
    try:
        data = path.read_bytes()
        text, encoding = decode_text_bytes(data)

        return build_single_page_document(
            path=path,
            text=text,
            method="txt",
            extra_warnings=[
                f"TXT file decoded using {encoding}."
            ],
        )

    except Exception as error:
        return failed_document(
            path,
            f"TXT extraction failed: {error}",
        )


def extract_xml(path: Path) -> ExtractedDocument:
    try:
        root = ElementTree.fromstring(
            path.read_bytes()
        )

        text_parts = [
            text.strip()
            for text in root.itertext()
            if text and text.strip()
        ]

        text = "\n".join(text_parts)

        return build_single_page_document(
            path=path,
            text=text,
            method="xml",
            extra_warnings=[
                "XML content was extracted from structured text nodes."
            ],
        )

    except Exception as error:
        return failed_document(
            path,
            f"XML extraction failed: {error}",
        )


def extract_document(
    document_path: str | Path,
) -> ExtractedDocument:
    path = Path(document_path)

    if not path.exists():
        return failed_document(
            path,
            f"File not found: {path}",
        )

    if not path.is_file():
        return failed_document(
            path,
            f"Expected a file but received: {path}",
        )

    extension = path.suffix.casefold()

    if extension == ".pdf":
        return extract_pdf(path)

    if extension == ".docx":
        return extract_docx(path)

    if extension == ".txt":
        return extract_txt(path)

    if extension == ".xml":
        return extract_xml(path)

    return failed_document(
        path,
        f"Unsupported file type: {extension or 'no extension'}",
    )


def main() -> None:
    parser = argparse.ArgumentParser(
        description=(
            "Extract resume text from PDF, DOCX, TXT, or XML."
        )
    )

    parser.add_argument(
        "document_path",
        help="Path to the resume file",
    )

    arguments = parser.parse_args()

    result = extract_document(arguments.document_path)
    print(result.model_dump_json(indent=2))


if __name__ == "__main__":
    main()